"""Normal, boundary and error checks for the expanded office workflow."""
from __future__ import annotations
import copy
import sys
import tempfile
import unittest
from datetime import date,timedelta
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"scripts"))
from skill_runtime import Runtime

class NewSkillTests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory()
        self.rt=Runtime(self.folder.name)
        self.demo=self.rt.add_file("monthly-finance-demo.xlsx",(ROOT/"outputs"/"atlas-finance-demo"/"monthly-finance-demo.xlsx").read_bytes())
    def tearDown(self): self.folder.cleanup()

    def test_search_and_reviewed_excel_edit_preserve_original(self):
        result=self.rt.submit("spreadsheet-search",{"file_id":self.demo["id"],"query":"Flights"})
        self.assertEqual(result["status"],"completed")
        self.assertEqual(result["output"]["matches"][0]["cell"],"H10")
        edit=self.rt.submit("spreadsheet-edit",{"file_id":self.demo["id"],"sheet":"BvA","cell":"J10","value":28500})
        self.assertEqual(edit["status"],"awaiting_approval")
        old=load_workbook(self.rt.file(self.demo["id"])[1],read_only=True)
        self.assertEqual(old["BvA"]["J10"].value,18500);old.close()
        done=self.rt.approve(edit["id"])
        self.assertEqual(done["status"],"completed",done["output"])
        new=load_workbook(self.rt.file(done["output"]["artifacts"][0]["id"])[1],read_only=True)
        self.assertEqual(new["BvA"]["J10"].value,28500);new.close()
        blocked=self.rt.submit("spreadsheet-edit",{"file_id":self.demo["id"],"sheet":"BvA","cell":"L10","value":1})
        self.assertEqual(self.rt.approve(blocked["id"])["status"],"needs_attention")

    def test_booking_finance_side_by_side_and_invoice_draft(self):
        booking=self.rt.submit("booking-confirm",{"customer":"Demo Client","product":"Chengdu launch","travel_date":"2026-10-25","quote_cny":80000,"supplier_option":"Sample option A"})
        self.assertEqual(booking["status"],"completed")
        record_id=booking["output"]["booking"]["id"]
        posting=self.rt.submit("finance-posting",{"booking_id":record_id,"file_id":self.demo["id"],"category":"Flights","committed_cny":22000})
        self.assertEqual(posting["status"],"awaiting_approval")
        done=self.rt.approve(posting["id"])
        self.assertEqual(done["status"],"completed",done["output"])
        self.assertEqual(done["output"]["new_pipeline_cny"],80000)
        self.assertEqual(done["output"]["new_committed_cny"],22000)
        updated=load_workbook(self.rt.file(done["output"]["artifacts"][0]["id"])[1],read_only=True)
        try:
            self.assertEqual(updated["BvA"]["A14"].value,"Demo Client")
            self.assertEqual(updated["BvA"]["D14"].value,80000)
            self.assertEqual(updated["BvA"]["K10"].value,22000)
            self.assertEqual(updated["BvA"]["C14"].value,0)
            self.assertEqual(updated["BvA"]["J10"].value,18500)
            self.assertEqual(updated["Booking log"]["A6"].value,record_id)
        finally:
            updated.close()
        duplicate=self.rt.submit("finance-posting",{"booking_id":record_id,"file_id":self.demo["id"],"category":"Flights","committed_cny":22000})
        self.assertEqual(self.rt.approve(duplicate["id"])["status"],"needs_attention")

    def test_general_business_packages(self):
        q=copy.deepcopy(self.rt.skill("quotation-package")["example"])
        q["valid_until"]=(date.today()+timedelta(days=30)).isoformat()
        result=self.rt.submit("quotation-package",q)
        self.assertEqual(result["status"],"completed",result["output"])
        self.assertEqual(result["output"]["total"],10070)
        self.assertEqual(len(result["output"]["artifacts"]),2)
        result=self.rt.submit("proposal-package",copy.deepcopy(self.rt.skill("proposal-package")["example"]))
        self.assertEqual(result["status"],"completed",result["output"])
        self.assertEqual(result["output"]["recommended"],"Pilot first")
        result=self.rt.submit("finance-forecast",copy.deepcopy(self.rt.skill("finance-forecast")["example"]))
        self.assertEqual(result["status"],"completed",result["output"])
        self.assertEqual(result["output"]["months"][0]["revenue"],66000)
        result=self.rt.submit("shareholder-report",copy.deepcopy(self.rt.skill("shareholder-report")["example"]))
        self.assertEqual(result["status"],"completed",result["output"])
        self.assertEqual(result["output"]["metrics"]["revenue_growth_pct"],12.5)

    def test_forecast_and_booking_invalid_inputs_are_visible(self):
        bad=copy.deepcopy(self.rt.skill("finance-forecast")["example"])
        bad["periods"][0]["churn"]=100
        self.assertEqual(self.rt.submit("finance-forecast",bad)["status"],"failed")
        bad=copy.deepcopy(self.rt.skill("booking-confirm")["example"])
        bad["quote_cny"]=0
        with self.assertRaises(ValueError): self.rt.submit("booking-confirm",bad)

if __name__=="__main__": unittest.main()
