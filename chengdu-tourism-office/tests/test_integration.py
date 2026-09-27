"""Integration checks for artifacts, local search and MCP tool wiring."""
from __future__ import annotations

import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pptx import Presentation
from mcp import Client

from scripts.tourism_core import demo
from scripts.render_reports import render
from scripts.search_store import index_case, search_cases
from server import tourism_mcp
from app import agent_reply, validate_company, simulate, update_case
from scripts.tourism_core import COMPANY
import copy


class IntegrationTests(unittest.TestCase):
    def test_programmable_people_and_avatar_reply(self):
        company = copy.deepcopy(COMPANY)
        validate_company(company)
        company["people"][0]["follows"] = ["missing_person"]
        with self.assertRaises(ValueError): validate_company(company)
        self.assertIn("configured skills", agent_reply({"person_id": "sales_manager", "question": "What do you do?"})["answer"])

    def test_option_change_invalidates_price_approvals(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch("app.DATA", Path(folder)), patch("app.COMPANY_FILE", Path(folder) / "company.json"):
                case = simulate({"message": "RFQ: 30 travelers for a 3 day Chengdu food and tea tour from 2026-10-25, budget CNY 200000."})
                self.assertIn("FINANCE", case["execution"]["data"]["completed"])
                alternative = case["comparison"]["data"]["alternatives"][0]["id"]
                changed = update_case({"request_id": case["request"]["request_id"], "option_id": alternative}, "option")
                self.assertEqual(changed["selected_option"]["id"], alternative)
                self.assertNotIn("FINANCE", changed["execution"]["data"]["completed"])
                self.assertIn("FINANCE", changed["invalidated_tasks"])

    def test_documents_and_search(self):
        case = demo()
        with tempfile.TemporaryDirectory() as folder:
            paths = render(case, folder)
            doc = Document(paths["docx"])
            slides = Presentation(paths["pptx"])
            self.assertTrue(any("Management decision" in p.text for p in doc.paragraphs))
            self.assertGreaterEqual(len(slides.slides), 7)
            db = Path(folder) / "cases.sqlite"
            index_case(db, case)
            self.assertEqual(search_cases(db, "tea")[0]["request_id"], case["request"]["request_id"])
            self.assertEqual(search_cases(db, "notfoundterm"), [])

    def test_mcp_tools_and_send_gate(self):
        async def run():
            async with Client(tourism_mcp.server, raise_exceptions=True) as client:
                listed = await client.list_tools()
                names = {tool.name for tool in listed.tools}
                self.assertTrue({"compare_flights", "find_tours", "search_prior_cases",
                                 "preview_report_email", "send_approved_report_email"}.issubset(names))
                result = await client.call_tool("compare_flights", {
                    "group_size": 30, "travel_start": "2026-10-25", "duration_days": 3})
                self.assertFalse(result.is_error)
                self.assertIn("FL-B", str(result.structured_content or result.content))
        asyncio.run(run())
        with tempfile.TemporaryDirectory() as folder:
            case = demo()
            tourism_mcp.OUTPUT = Path(folder)
            (Path(folder) / "SIM-RFQ-001-trace.json").write_text(json.dumps(case), encoding="utf-8")
            render(case, folder)
            with self.assertRaises(PermissionError):
                tourism_mcp.send_approved_report_email("SIM-RFQ-001", "not-approved")


if __name__ == "__main__":
    unittest.main()
