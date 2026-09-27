"""Auditable local file and simulated-booking operations."""
from __future__ import annotations

import io
import math
import re
import zipfile
from datetime import datetime, timezone

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties


def _workbook(rt, file_id):
    meta, path = rt.file(file_id)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Choose an Excel .xlsx workspace file.")
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 80_000_000:
            raise ValueError("Expanded workbook is too large.")
    book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    if len(book.worksheets) > 25 or any(s.max_row > 20_000 or s.max_column > 100 for s in book):
        book.close()
        raise ValueError("Workbooks are limited to 25 sheets, 20,000 rows and 100 columns per sheet.")
    return meta, book


def search_workbook(rt, payload):
    meta, book = _workbook(rt, payload["file_id"])
    query = payload["query"].casefold().strip()
    if len(query) < 2:
        raise ValueError("Search for at least two characters.")
    matches = []
    try:
        for sheet in book:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None and query in str(cell.value).casefold():
                        matches.append({"sheet": sheet.title, "cell": cell.coordinate,
                                        "value": str(cell.value)[:300]})
                        if len(matches) >= 100: break
                if len(matches) >= 100: break
            if len(matches) >= 100: break
    finally:
        book.close()
    return {"summary": f"Found {len(matches)} matching cells in {meta['name']}.",
            "matches": matches, "truncated": len(matches) >= 100}


def edit_workbook(rt, payload):
    meta, book = _workbook(rt, payload["file_id"])
    try:
        if payload["sheet"] not in book.sheetnames:
            raise ValueError("The selected worksheet does not exist.")
        if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,4}", payload["cell"].upper()):
            raise ValueError("Use one exact Excel cell, such as C10.")
        cell = book[payload["sheet"]][payload["cell"].upper()]
        if cell.data_type == "f":
            raise ValueError("Formula cells cannot be overwritten; edit an input cell instead.")
        value = payload["value"]
        if isinstance(value, (dict, list, bool)) or (isinstance(value, float) and not math.isfinite(value)):
            raise ValueError("Use a finite number or plain text value.")
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            raise ValueError("Formula-like text is blocked. Use a plain value.")
        before = cell.value
        cell.value = value
        book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
        buf = io.BytesIO(); book.save(buf)
    finally:
        book.close()
    output = rt.add_file(meta["name"].removesuffix(".xlsx") + "-edited.xlsx", buf.getvalue())
    return {"summary": "A reviewed input cell was changed in a new workbook copy. Excel recalculates formulas when opened.",
            "change": {"sheet": payload["sheet"], "cell": cell.coordinate, "before": before, "after": value},
            "artifacts": [output], "original_file_id": payload["file_id"]}


def simulated_booking(rt, payload):
    if payload["quote_cny"] <= 0:
        raise ValueError("A simulated booking needs a positive quoted amount.")
    record = rt.record("booking-confirm", "Simulated booking: " + payload["customer"],
                       {**payload, "status": "simulated_confirmed", "real_booking": False})
    return {"summary": "Simulated confirmation recorded. No supplier reservation, payment or customer email occurred.",
            "booking": record, "next_step": "Finance can prepare a committed-budget update and invoice draft for director review."}


def finance_posting(rt, payload):
    booking = next((r for r in rt.records("booking-confirm") if r["id"] == payload["booking_id"]), None)
    if not booking or booking["payload"].get("status") != "simulated_confirmed":
        raise ValueError("Choose an existing simulated confirmation record.")
    if any(r["payload"].get("booking_id") == payload["booking_id"] for r in rt.records("finance-posting")):
        raise ValueError("This booking already has a reviewed finance posting.")
    if payload["committed_cny"] <= 0 or payload["committed_cny"] > booking["payload"]["quote_cny"]:
        raise ValueError("Committed cost must be positive and cannot exceed the quoted customer amount.")
    meta, book = _workbook(rt, payload["file_id"])
    try:
        if "BvA" not in book or "Booking log" not in book:
            raise ValueError("Use the Atlas monthly finance demo workbook with BvA and Booking log sheets.")
        sheet, log = book["BvA"], book["Booking log"]
        category = payload["category"]
        expense_row = next((n for n in range(10, 30) if sheet[f"H{n}"].value == category), None)
        if expense_row is None:
            raise ValueError("Choose a category shown on the BvA sheet.")
        customer = booking["payload"]["customer"]
        income_row = next((n for n in range(10, 30) if sheet[f"A{n}"].value == customer), None)
        if income_row is None:
            income_row = next((n for n in range(10, 30) if sheet[f"A{n}"].value in (None, "")), None)
            if income_row is None:
                raise ValueError("The income section has no empty customer row.")
            sheet[f"A{income_row}"] = customer
            sheet[f"B{income_row}"] = 0
            sheet[f"C{income_row}"] = 0
            sheet[f"D{income_row}"] = 0
        previous_income = sheet[f"D{income_row}"].value
        previous = sheet[f"K{expense_row}"].value
        if not isinstance(previous, (int, float)):
            raise ValueError("Committed budget input must be numeric.")
        if not isinstance(previous_income, (int, float)):
            raise ValueError("Simulated income pipeline input must be numeric.")
        new_committed = round(previous + payload["committed_cny"], 2)
        sheet[f"K{expense_row}"] = new_committed
        new_income = round(previous_income + booking["payload"]["quote_cny"], 2)
        sheet[f"D{income_row}"] = new_income
        log_row = next((n for n in range(6, log.max_row + 2) if log[f"A{n}"].value in (None, "")), log.max_row + 1)
        log_values = [payload["booking_id"], booking["payload"]["customer"], "simulated_confirmed",
                      booking["payload"]["quote_cny"], category, payload["committed_cny"],
                      "Reviewed in Activity", datetime.now(timezone.utc).isoformat()]
        for col, value in enumerate(log_values, 1):
            log.cell(log_row, col, value)
        book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
        buf = io.BytesIO(); book.save(buf)
    finally:
        book.close()
    finance = rt.add_file(meta["name"].removesuffix(".xlsx") + "-reviewed.xlsx", buf.getvalue())
    from business_tools import workbook
    invoice = workbook(rt, "DRAFT invoice - " + booking["payload"]["customer"],
                       ["Item", "Amount (CNY)", "Status"],
                       [["Tourism product launch package", booking["payload"]["quote_cny"], "DRAFT - NOT ISSUED"]])
    record = rt.record("finance-posting", "Reviewed finance posting: " + booking["payload"]["customer"],
                       {**payload, "status": "reviewed_simulated_commitment", "finance_file_id": finance["id"],
                        "invoice_file_id": invoice["id"]})
    return {"summary": "Reviewed simulated customer income and supplier commitment were added to projections in a new workbook; invoice draft created. Actuals did not change. No payment or customer invoice was issued.",
            "previous_committed_cny": previous, "new_committed_cny": new_committed,
            "previous_pipeline_cny": previous_income, "new_pipeline_cny": new_income,
            "record_id": record["id"], "artifacts": [finance, invoice]}
