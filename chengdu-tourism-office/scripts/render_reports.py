"""Render the simulated RFQ outcome as reviewable Word and PowerPoint files."""
from __future__ import annotations

from pathlib import Path


def render(case: dict, output_dir: str | Path) -> dict[str, str]:
    from docx import Document
    from pptx import Presentation

    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    request_id = case["request"]["request_id"]
    report = case["outcome"]["data"]
    inquiry = case["inquiry"]["data"]
    comparison = case["comparison"]["data"]
    schedule = case["schedule"]["data"]
    team = case["hierarchy"]["data"]
    work = case["workload"]["data"]

    document = Document()
    document.add_heading("Tourism RFQ | Management decision brief", 0)
    document.add_paragraph(f"Request {request_id} · SIMULATION · independent demonstration")
    document.add_heading("Customer request", 1)
    document.add_paragraph(inquiry["message"])
    document.add_paragraph(f"{inquiry['group_size']} travelers · {inquiry['duration_days']} days · budget CNY {inquiry['budget_cny']}")
    document.add_heading("Options and price composition", 1)
    table = document.add_table(rows=1, cols=5)
    for cell, label in zip(table.rows[0].cells, ("Option", "Land", "Flight", "Activities", "Total CNY")): cell.text = label
    for option in [comparison["recommended"], *comparison["alternatives"]]:
        row = table.add_row().cells
        for cell, value in zip(row, (option["id"], option["land_quote_cny"], option["flight_estimate_cny"],
                                     option["activity_estimate_cny"], option["quote_cny"])): cell.text = str(value)
    document.add_paragraph("All offers are indicative. Flight fares, seats, supplier inventory and final terms require confirmation.")
    document.add_heading("Itinerary schedule", 1)
    for day in schedule["days"]:
        document.add_heading(f"Day {day['day_number']} · {day['date']}", 2)
        for item in day["activities"]: document.add_paragraph(f"{item['slot'].title()}: {item['title']}", style="List Bullet")
    document.add_heading("Organizational handoffs", 1)
    document.add_paragraph(f"Accountable task lead: {team['task_lead']}")
    for task in work["tasks"]:
        document.add_paragraph(f"{task['id']} → {task['owner_id']} · due {task['due_date']} · route {' → '.join(task['route_from_lead'])}", style="List Bullet")
    document.add_heading("Control and outcome", 1)
    document.add_paragraph(report["management_summary"])
    document.add_paragraph(f"Decision: {report['decision']}; email: draft only.")
    for task_id, evidence in report["evidence_register"].items():
        document.add_paragraph(f"{task_id}: {evidence}", style="List Bullet")
    document.add_heading("Source and limitation", 1)
    document.add_paragraph("Company hierarchy, fares, tours, task evidence and financial assumptions are synthetic demonstration data. Public company materials provide context only; no affiliation or actual quotation is claimed.")
    docx_path = output / f"{request_id}-decision-brief.docx"
    document.save(docx_path)

    presentation = Presentation()
    def slide(title: str, lines: list[str]) -> None:
        page = presentation.slides.add_slide(presentation.slide_layouts[1])
        page.shapes.title.text = title
        page.placeholders[1].text = "\n".join(lines)
    slide("Tourism RFQ | simulated company workflow", [request_id, inquiry["message"], "Independent demonstration"])
    slide("Commercial options", [f"{x['id']}: CNY {x['quote_cny']} total ({x['quote_per_person_cny']} per traveler)"
                                  for x in [comparison["recommended"], *comparison["alternatives"]]])
    slide("People and delegation", [f"Task lead: {team['task_lead']}",
                                     *[f"{x['id']} → {x['owner_id']}" for x in work["tasks"][:7]]])
    for day in schedule["days"]:
        slide(f"Day {day['day_number']} | {day['date']}", [f"{x['slot']}: {x['title']}" for x in day["activities"]])
    slide("Decision and controls", [report["management_summary"], f"Decision: {report['decision']}",
                                    f"Evidence records: {len(report['evidence_register'])}",
                                    "Customer email remains a draft until separately approved."])
    slide("Evidence limits", ["Synthetic company roster and prices", "Supplier and flight inventory unconfirmed",
                              "No real email sent or booking placed"])
    pptx_path = output / f"{request_id}-management-deck.pptx"
    presentation.save(pptx_path)
    return {"docx": str(docx_path), "pptx": str(pptx_path)}
