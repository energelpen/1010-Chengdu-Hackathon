"""Deterministic business tools. No model output is executed as code."""
from __future__ import annotations
import csv
import io
import json
import math
import re
import statistics
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from email.message import EmailMessage
from pathlib import Path
from xml.sax.saxutils import escape

def cash(n): return float(Decimal(str(n)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP))
def parse_date(s):
    try: return date.fromisoformat(s)
    except (ValueError,TypeError): raise ValueError("Use a valid date in YYYY-MM-DD format.")
def time_range(p):
    try: start,end=(datetime.fromisoformat(p[k].replace("Z","+00:00")) for k in ("start","end"))
    except (ValueError,TypeError): raise ValueError("Use ISO date/time values with an explicit timezone offset.")
    if start.tzinfo is None or end.tzinfo is None: raise ValueError("Both dates need a timezone offset, such as +08:00.")
    if end<=start: raise ValueError("End must be after start.")
    return start,end
def email_message(p):
    if not re.fullmatch(r"[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+",p["to"]): raise ValueError("Enter one valid recipient email address.")
    if any(c in p["subject"] for c in "\r\n"): raise ValueError("Subject must be a single line.")
    m=EmailMessage(); m["To"]=p["to"]; m["Subject"]=p["subject"]; m.set_content(p["body"])
    return m
def label(k): return k.replace("_"," ").capitalize()
def markdown(p):
    lines=["# "+str(p.get("title",p.get("employee","Company brief"))), ""]
    for key,val in p.items():
        if key=="title": continue
        lines += ["## "+label(key), ""]
        if isinstance(val,list):
            for item in val:
                lines.append("- "+("; ".join(f"{label(k)}: {v}" for k,v in item.items()) if isinstance(item,dict) else str(item)))
        else: lines.append(str(val))
        lines.append("")
    return "\n".join(lines)
def workbook(rt,title,columns,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table,TableStyleInfo
    if not columns or len(set(map(str,columns)))!=len(columns): raise ValueError("Use distinct, nonempty column headers.")
    if any(len(row)!=len(columns) for row in rows): raise ValueError("Every row must have the same number of cells as the columns.")
    wb=Workbook(); ws=wb.active; ws.title="Overview"
    ws.append(columns)
    for row in rows: ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type="s"
            cell.alignment=Alignment(vertical="top",wrap_text=True)
    for c in ws[1]: c.fill=PatternFill("solid",fgColor="246B5A"); c.font=Font(color="FFFFFF",bold=True)
    for i,col in enumerate(columns,1): ws.column_dimensions[get_column_letter(i)].width=min(44,max(18,len(str(col))+5))
    ws.freeze_panes="A2"
    if rows:
        table=Table(displayName="AtlasData",ref=f"A1:{get_column_letter(len(columns))}{len(rows)+1}")
        table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        ws.add_table(table)
    buf=io.BytesIO(); wb.save(buf)
    return rt.add_file(title[:100]+".xlsx",buf.getvalue())
def read_table(rt,id):
    meta,path=rt.file(id)
    if path.suffix==".xlsx":
        from openpyxl import load_workbook
        import zipfile
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>80_000_000: raise ValueError("Expanded workbook is too large.")
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            ws=wb.active
            if ws.max_row and ws.max_row>20000 or ws.max_column and ws.max_column>100: raise ValueError("Analyze up to 20,000 rows and 100 columns.")
            rows=[]
            for i,row in enumerate(ws.iter_rows(values_only=True)):
                if i>20000 or len(row)>100: raise ValueError("Analyze up to 20,000 rows and 100 columns.")
                rows.append(list(row))
        finally: wb.close()
    elif path.suffix==".csv":
        try: rows=list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
        except UnicodeError: raise ValueError("Upload a UTF-8 CSV file.")
        if len(rows)>20001 or any(len(r)>100 for r in rows): raise ValueError("Analyze up to 20,000 rows and 100 columns.")
    else: raise ValueError("Choose an Excel .xlsx or UTF-8 .csv file.")
    if not rows: raise ValueError("The file has no rows.")
    if not rows[0] or any(len(r)!=len(rows[0]) for r in rows): raise ValueError("Every row must match the header width.")
    return rows[0],rows[1:]
def profile(columns,rows):
    result=[]
    for i,col in enumerate(columns):
        vals=[r[i] for r in rows]; nums=[]
        nonempty=[v for v in vals if v not in (None,"")]
        for v in nonempty:
            try:
                n=float(v)
                if not math.isfinite(n) or isinstance(v,bool): continue
                nums.append(n)
            except (ValueError,TypeError): pass
        result.append({"column":str(col),"missing":len(vals)-len(nonempty),"numeric_count":len(nums),
            "sum":cash(sum(nums)) if nums else None,"mean":cash(statistics.mean(nums)) if nums else None,
            "min":min(nums) if nums else None,"max":max(nums) if nums else None})
    return {"row_count":len(rows),"columns":result,"note":"Excel formula values use the file's saved cache; formulas are not recalculated."}
def document(rt,p):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.8)
    style=doc.styles["Normal"]; style.font.name="Calibri"; style.font.size=Pt(11)
    doc.add_heading(p["title"],0)
    for section in p["sections"]:
        doc.add_heading(section["heading"],1)
        for para in section["body"].split("\n"): doc.add_paragraph(para)
    doc.core_properties.title=p["title"]
    buf=io.BytesIO(); doc.save(buf)
    return rt.add_file(p["title"][:100]+".docx",buf.getvalue())
def presentation(rt,p):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    for i,item in enumerate([{"title":p["title"],"bullets":["Prepared with Atlas · Company workspace"]}]+p["slides"]):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string("F6F8F5")
        def box(x,y,w,h,text,size,color):
            shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
            frame=shape.text_frame; frame.word_wrap=True
            for j,line in enumerate(text.split("\n")):
                para=frame.paragraphs[0] if j==0 else frame.add_paragraph()
                para.text=line; para.font.size=Pt(size); para.font.name="Aptos"; para.font.color.rgb=RGBColor.from_string(color); para.space_after=Pt(15)
            return shape
        if len(item["title"])>120: raise ValueError("Keep slide titles under 120 characters.")
        box(.8,.4,11,.3,"ATLAS / COMPANY WORKSPACE",11,"487B6C")
        box(.8,1.1,11.7,1.5,item["title"],34 if len(item["title"])<65 else 28,"193F35")
        box(.9,2.9,11.3,3.8,"\n".join("• "+x for x in item["bullets"]),22 if sum(map(len,item["bullets"]))<500 else 18,"394C46")
        box(11.7,7,1,.25,f"{i+1:02}",10,"487B6C")
    buf=io.BytesIO(); prs.save(buf)
    return rt.add_file(p["title"][:100]+".pptx",buf.getvalue())
def pdf(rt,p):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    styles=getSampleStyleSheet()
    font=Path("C:/Windows/Fonts/arial.ttf")
    if font.exists():
        if "AtlasArial" not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont("AtlasArial",str(font)))
        for style in styles.byName.values(): style.fontName="AtlasArial"
    story=[Paragraph(escape(p["title"]),styles["Title"]),Spacer(1,20)]
    for s in p["sections"]:
        story.extend([Paragraph(escape(s["heading"]),styles["Heading2"]),Paragraph(escape(s["body"]).replace("\n","<br/>"),styles["BodyText"]),Spacer(1,12)])
    buf=io.BytesIO(); SimpleDocTemplate(buf,title=p["title"]).build(story)
    return rt.add_file(p["title"][:100]+".pdf",buf.getvalue())
def roster(rt):
    from skill_runtime import ROOT
    path=rt.data/"company.json"
    company=json.loads((path if path.exists() else ROOT/"resources"/"company-general.json").read_text(encoding="utf-8"))
    return company["people"]
def resolve_person(people,value):
    match=next((x for x in people if x["id"].casefold()==value.casefold() or x["name"].casefold()==value.casefold()),None)
    if not match: raise ValueError("Unknown colleague: "+value+". Choose a person from the current org chart.")
    return match
def execute(rt,skill,p,actor="atlas"):
    kind=skill["handler"]; sid=skill["id"]
    if kind in ("quotation_package","proposal_package","finance_forecast","shareholder_package"):
        from business_packages import quotation,proposal,forecast,shareholder
        return {"quotation_package":quotation,"proposal_package":proposal,
                "finance_forecast":forecast,"shareholder_package":shareholder}[kind](rt,p)
    if kind=="telegram":
        from telegram_service import notify
        return notify(rt,p,actor)
    if kind in ("workbook_search", "workbook_edit", "simulated_booking", "finance_posting"):
        from office_operations import search_workbook, edit_workbook, simulated_booking, finance_posting
        return {"workbook_search": search_workbook, "workbook_edit": edit_workbook,
                "simulated_booking": simulated_booking, "finance_posting": finance_posting}[kind](rt, p)
    if kind=="google":
        from google_workspace import GoogleWorkspace
        return GoogleWorkspace(rt.data).execute(sid,p,rt,actor=actor)
    if kind=="mcp":
        from mcp_bridge import Bridge
        return Bridge(rt.data).call(p)
    if kind=="tourism":
        import tourism_core
        return getattr(tourism_core,skill["function"])(p)
    if kind=="email":
        f=rt.add_file("email-draft.eml",email_message(p).as_bytes())
        return {"summary":"Email draft created. It has not been sent.","artifacts":[f],"delivery":"draft"}
    if kind=="spreadsheet": return {"summary":"Excel workbook created.","artifacts":[workbook(rt,p["title"],p["columns"],p["rows"])]}
    if kind=="analyze": return {"summary":"Spreadsheet profile complete.","data":profile(*read_table(rt,p["file_id"]))}
    if kind=="document": return {"summary":"Word document created.","artifacts":[document(rt,p)]}
    if kind=="presentation": return {"summary":"PowerPoint presentation created.","artifacts":[presentation(rt,p)]}
    if kind=="pdf": return {"summary":"PDF brief created.","artifacts":[pdf(rt,p)]}
    if kind=="extract_pdf":
        from pypdf import PdfReader
        _,path=rt.file(p["file_id"])
        if path.suffix!=".pdf": raise ValueError("Choose a PDF file.")
        reader=PdfReader(path)
        if len(reader.pages)>200: raise ValueError("Read up to 200 PDF pages at a time.")
        pages=[{"page":i+1,"text":page.extract_text()[:30000]} for i,page in enumerate(reader.pages)]
        return {"summary":"PDF text extracted; image-only pages need OCR.","pages":pages}
    if kind=="clean_csv":
        _,path=rt.file(p["file_id"])
        if path.suffix!=".csv": raise ValueError("Choose a CSV file.")
        cols,rows=read_table(rt,p["file_id"]); clean=[[str(v).strip() for v in r] for r in rows]
        if p["deduplicate"]: clean=[list(r) for r in dict.fromkeys(tuple(r) for r in clean)]
        buf=io.StringIO(newline=""); writer=csv.writer(buf)
        # CSV cannot represent a text cell type; prefix risky leading characters for office applications.
        def safe_row(row): return [("'"+v if isinstance(v,str) and v.startswith(("=","+","-","@","\t","\r")) else v) for v in row]
        writer.writerow(safe_row(cols)); writer.writerows(safe_row(r) for r in clean)
        return {"summary":f"Cleaned {len(rows)} rows into {len(clean)} rows. Formula-like cells are escaped for spreadsheet safety.","data":profile(cols,clean),"artifacts":[rt.add_file("cleaned.csv",buf.getvalue().encode("utf-8-sig"))]}
    if kind=="calendar":
        start,end=time_range(p)
        def esc(v): return v.replace("\\","\\\\").replace("\r","").replace("\n","\\n").replace(";","\\;").replace(",","\\,")
        lines=["BEGIN:VCALENDAR","VERSION:2.0","PRODID:-//Atlas//Company Workspace//EN","BEGIN:VEVENT","UID:"+uuid.uuid4().hex+"@atlas.local","DTSTAMP:"+datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),"DTSTART:"+start.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),"DTEND:"+end.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),"SUMMARY:"+esc(p["title"]),"DESCRIPTION:"+esc(p["description"]),"END:VEVENT","END:VCALENDAR"]
        folded=[]
        for line in lines:
            part=""
            for char in line:
                if len((part+char).encode("utf-8"))>73: folded.append(part); part=" "+char
                else: part+=char
            folded.append(part)
        return {"summary":"Calendar file created. No invitations sent.","artifacts":[rt.add_file("event.ics",("\r\n".join(folded)+"\r\n").encode())]}
    if kind=="project":
        tasks={t["id"]:t for t in p["tasks"]}; start=parse_date(p["start"])
        if len(tasks)!=len(p["tasks"]): raise ValueError("Task IDs must be unique.")
        finished={}; visiting=set()
        def visit(id):
            if id not in tasks: raise ValueError("Unknown dependency: "+id)
            if id in visiting: raise ValueError("Project dependencies contain a cycle.")
            if id in finished: return finished[id]
            visiting.add(id); task=tasks[id]
            begins=max([start]+[visit(dep)["finish"] for dep in task["depends_on"]])
            visiting.remove(id); finished[id]={**task,"start":begins,"finish":begins+timedelta(days=task["days"])}
            return finished[id]
        for id in tasks: visit(id)
        data=[{**t,"start":t["start"].isoformat(),"finish":t["finish"].isoformat()} for t in finished.values()]
        return {"summary":"Schedule calculated in calendar days; finish dates are exclusive.","tasks":data,"artifacts":[workbook(rt,p["title"],["Task","Owner","Start","Finish","Days"],[[t["title"],t["owner"],t["start"],t["finish"],t["days"]] for t in data])]}
    if kind in ("record","knowledge"):
        result=rt.record(sid,p.get("title",sid),p)
        return {"summary":"Saved to the local company register.","record":result}
    if kind=="knowledge_search": return {"summary":"Matching local source notes.","matches":rt.records("knowledge-save",p["query"])[:30]}
    if kind=="pipeline":
        stages=defaultdict(float)
        for d in p["deals"]: stages[d["stage"]]+=d["value"]*d["probability"]
        return {"summary":"Weighted pipeline calculated.","currency":p["currency"],"total":cash(sum(d["value"] for d in p["deals"])),"weighted":cash(sum(stages.values())),"by_stage":{k:cash(v) for k,v in stages.items()}}
    if kind=="invoice":
        items=[{**x,"total":cash(Decimal(str(x["quantity"]))*Decimal(str(x["unit_price"])))} for x in p["items"]]
        subtotal=cash(sum(Decimal(str(x["total"])) for x in items)); tax=cash(Decimal(str(subtotal))*Decimal(str(p["tax_rate"])))
        total=cash(Decimal(str(subtotal))+Decimal(str(tax)))
        rows=[[x["description"],x["quantity"],x["unit_price"],x["total"]] for x in items]+[["Subtotal",None,None,subtotal],["Tax",None,None,tax],["Total",None,None,total]]
        return {"summary":"Invoice draft calculated; not issued or posted.","currency":p["currency"],"subtotal":subtotal,"tax":tax,"total":total,"artifacts":[workbook(rt,"Invoice - "+p["customer"],["Description","Quantity","Unit price ("+p["currency"]+")","Total ("+p["currency"]+")"],rows)]}
    if kind=="expenses":
        groups=defaultdict(float)
        for x in p["expenses"]: groups[x["category"]]+=x["amount"]
        return {"summary":"Expenses prepared for review.","currency":p["currency"],"total":cash(sum(groups.values())),"categories":{k:cash(v) for k,v in groups.items()},"missing_receipts":[x["description"] for x in p["expenses"] if not x["receipt"]]}
    if kind=="variance":
        return {"summary":"Positive variance means overspending.","currency":p["currency"],"items":[{**x,"variance":cash(x["actual"]-x["budget"]),"variance_percent":round(100*(x["actual"]-x["budget"])/x["budget"],2) if x["budget"] else None} for x in p["items"]]}
    if kind=="procurement":
        weights=[p[x+"_weight"] for x in ("cost","quality","delivery")]
        if sum(weights)<=0 or not p["suppliers"]: raise ValueError("Supply candidates and at least one positive weight.")
        maxcost=max(x["cost"] for x in p["suppliers"]); maxdays=max(x["delivery_days"] for x in p["suppliers"])
        result=[]
        for x in p["suppliers"]:
            components=[1-x["cost"]/maxcost if maxcost else 1,x["quality"]/100,1-x["delivery_days"]/maxdays if maxdays else 1]
            result.append({**x,"score":round(100*sum(a*b for a,b in zip(weights,components))/sum(weights),2),"components":dict(zip(["cost","quality","delivery"],components))})
        return {"summary":"Suppliers ranked using your weights; no order placed.","ranking":sorted(result,key=lambda x:-x["score"])}
    if kind=="inventory":
        result=[]
        for x in p["items"]:
            point=x["daily_demand"]*x["lead_days"]+x["safety_stock"]; position=x["on_hand"]+x["on_order"]
            target=max(point,x["daily_demand"]*x["target_days"]+x["safety_stock"])
            result.append({**x,"reorder_point":point,"inventory_position":position,"suggested_quantity":max(0,math.ceil(target-position)) if position<=point else 0})
        return {"summary":"Reorder recommendations calculated; no purchases placed.","items":result}
    if kind=="onboarding":
        start=parse_date(p["start"])
        tasks=[{"task":task,"owner":owner,"due":(start+timedelta(days=offset)).isoformat()} for task,owner,offset in [("Prepare equipment and account access","IT",-3),("Confirm role and first-week plan",p["manager"],0),("Complete policy orientation","People",0),("Review first-week progress",p["manager"],7),("Review first-month goals",p["manager"],30)]]
        return {"summary":"Onboarding checklist prepared.","actions":tasks,"artifacts":[rt.add_file("onboarding.md",markdown({**p,"title":"Onboarding: "+p["employee"],"actions":tasks}).encode())]}
    if kind=="scorecard":
        weight=sum(x["weight"] for x in p["criteria"])
        if weight<=0: raise ValueError("At least one criterion needs a positive weight.")
        return {"summary":"Scorecard calculated for human review; no hiring decision made.","candidate":p["candidate"],"score_out_of_5":round(sum(x["weight"]*x["score"] for x in p["criteria"])/weight,2),"criteria":p["criteria"]}
    if kind=="leave":
        start,end=parse_date(p["start"]),parse_date(p["end"])
        if end<start or (end-start).days>365: raise ValueError("Use an end date within one year after the start.")
        days=sum((start+timedelta(days=i)).weekday()<5 for i in range((end-start).days+1))
        return {"summary":"Leave request recorded; weekends excluded, holidays not considered.","weekdays":days,"status":"pending_manager_review","record":rt.record(sid,"Leave: "+p["employee"],p)}
    if kind=="triage":
        score={"low":1,"medium":2,"high":3}[p["impact"]]*{"low":1,"medium":2,"high":3}[p["urgency"]]
        priority="P1" if score>=6 else "P2" if score>=3 else "P3"
        return {"summary":"Ticket recorded with suggested priority.","priority":priority,"suggested_response_hours":{"P1":1,"P2":4,"P3":24}[priority],"record":rt.record(sid,p["title"],p)}
    if kind=="risks": return {"summary":"Risks ranked by likelihood × impact.","risks":sorted([{**x,"score":x["likelihood"]*x["impact"]} for x in p["risks"]],key=lambda x:-x["score"])}
    if kind=="checklist": return {"summary":"Checklist reviewed against supplied evidence.","counts":{s:sum(x["status"]==s for x in p["items"]) for s in ["met","gap","unknown"]},"artifacts":[rt.add_file(p["title"][:100]+".md",markdown(p).encode())]}
    if kind=="capacity":
        return {"summary":"Configured team capacity; no live calendar data.","people":[{"id":x["id"],"name":x["name"],"free_hours":x["capacity_hours"]-x["assigned_hours"] if x["available"] else 0,"skills":x["skills"]} for x in roster(rt)]}
    if kind=="raci":
        people=roster(rt); rows=[]; seen=set()
        for item in p["assignments"]:
            task=item["task"].strip().casefold()
            if task in seen: raise ValueError("Each task needs one responsibility row.")
            seen.add(task)
            row={"task":item["task"]}
            for field in ("responsible","consulted","informed"):
                row[field]=[resolve_person(people,value)["name"] for value in item[field]]
            row["accountable"]=resolve_person(people,item["accountable"])["name"]
            if not row["responsible"]: raise ValueError("Each task needs at least one responsible colleague.")
            rows.append(row)
        artifact=workbook(rt,p["project"]+" responsibility matrix",["Task","Responsible","Accountable","Consulted","Informed"],[[r["task"],", ".join(r["responsible"]),r["accountable"],", ".join(r["consulted"]),", ".join(r["informed"])] for r in rows])
        return {"summary":"Responsibility matrix prepared from the current roster; assignments are proposals.","rows":rows,"artifacts":[artifact]}
    if kind=="rebalance":
        people=roster(rt); used={x["id"]:x["assigned_hours"] for x in people}; assignments=[]; unassigned=[]
        for task in p["tasks"]:
            hours=task["hours"]; candidates=[x for x in people if x["available"] and task["required_skill"] in x["skills"] and x["capacity_hours"]-used[x["id"]]>=hours]
            preferred=task.get("preferred_person","").strip()
            if preferred:
                preferred_person=resolve_person(people,preferred)
                chosen=preferred_person if preferred_person in candidates else None
            else: chosen=min(candidates,key=lambda x:((used[x["id"]]+hours)/x["capacity_hours"],-x["capacity_hours"]+used[x["id"]],x["id"])) if candidates else None
            if not chosen:
                unassigned.append({"task":task["task"],"required_skill":task["required_skill"],"hours":hours,"reason":"Preferred colleague cannot take this work." if preferred else "No available colleague has the skill and remaining capacity."})
                continue
            used[chosen["id"]]+=hours
            assignments.append({"task":task["task"],"person_id":chosen["id"],"person":chosen["name"],"manager":next((x["name"] for x in people if x["id"]==chosen.get("reports_to")),"Company board"),"hours":hours,"projected_hours":used[chosen["id"]],"capacity_hours":chosen["capacity_hours"]})
        return {"summary":"Projected assignments prepared. No staff records were changed.","assignments":assignments,"unassigned":unassigned,"needs_manager_attention":bool(unassigned)}
    if kind=="scenario":
        criteria=p["criteria"]; options=p["options"]; weights=sum(x["weight"] for x in criteria)
        if weights<=0: raise ValueError("At least one criterion needs a positive weight.")
        if len({x["name"].casefold() for x in criteria})!=len(criteria): raise ValueError("Criterion names must be unique.")
        if len({x["name"].casefold() for x in options})!=len(options): raise ValueError("Option names must be unique.")
        ranking=[]
        for option in options:
            components=[]
            for criterion in criteria:
                name=criterion["name"]
                if name not in option["measures"]: raise ValueError("Missing "+name+" for "+option["name"]+".")
                values=[float(x["measures"][name]) for x in options]
                low,high=min(values),max(values)
                normalized=1 if low==high else ((float(option["measures"][name])-low)/(high-low) if criterion["direction"]=="higher" else (high-float(option["measures"][name]))/(high-low))
                components.append({"criterion":name,"value":option["measures"][name],"weighted_points":round(100*criterion["weight"]*normalized/weights,2)})
            ranking.append({"option":option["name"],"score":round(sum(x["weighted_points"] for x in components),2),"components":components})
        ranking.sort(key=lambda x:(-x["score"],x["option"].casefold()))
        return {"summary":"Options ranked from supplied measures; tied measures receive equal credit. Review assumptions before deciding.","ranking":ranking,"recommended":ranking[0]["option"],"weights_total":weights}
    if kind=="handoff":
        for item in p["items"]: parse_date(item["due"])
        status_counts={s:sum(item["status"]==s for item in p["items"]) for s in ("not_started","in_progress","blocked","done")}
        blocked=[{"work":item["work"],"owner":item["owner"],"blocker":item["blocker"]} for item in p["items"] if item["status"]=="blocked"]
        if any(item["status"]=="blocked" and not item["blocker"].strip() for item in p["items"]): raise ValueError("Describe the blocker for each blocked item.")
        sections=[{"heading":"At a glance","body":f"Audience: {p['audience']}\n"+", ".join(f"{status.replace('_',' ')}: {count}" for status,count in status_counts.items())},
                  {"heading":"Work and ownership","body":"\n".join(f"{x['work']} — {x['status'].replace('_',' ')}; owner: {x['owner']}; due: {x['due']}; next: {x['next_step']}"+(f"; blocker: {x['blocker']}" if x["blocker"] else "") for x in p["items"])},
                  {"heading":"Decisions needed","body":"\n".join(f"{x['work']}: {x['blocker']} — ask {x['owner']} for a resolution" for x in blocked) or "No reported blockers."}]
        artifact=document(rt,{"title":p["title"],"sections":sections})
        return {"summary":"Management handoff prepared from supplied status data; no messages sent.","counts":status_counts,"blockers":blocked,"artifacts":[artifact],"email_draft":{"subject":p["title"],"body":sections[0]["body"]+"\n\n"+sections[2]["body"]}}
    if kind=="brief": return {"summary":"Structured brief prepared from your supplied information.","artifacts":[rt.add_file(p["title"][:100]+".md",markdown(p).encode())]}
    raise ValueError("No registered implementation for this skill.")
