"""Record authentic Atlas UI operations in a fresh, credential-free workspace."""
import json, os, shutil, subprocess, time
from pathlib import Path
from urllib.request import urlopen
from playwright.sync_api import sync_playwright

BASE = Path(__file__).resolve().parents[2]
SRC = BASE/'chengdu-tourism-office'
QA = BASE/'submission/qa/full-demo'
APP = QA/'app'
DATA = QA/('state-'+str(int(time.time())))
QA.mkdir(parents=True, exist_ok=True)
for name in ('scripts','web','resources','shared','skills','examples'):
    shutil.copytree(SRC/name, APP/name, dirs_exist_ok=True, ignore=shutil.ignore_patterns('__pycache__','*.pyc','.env'))
shutil.copy2(SRC/'app.py', APP/'app.py')
env=dict(os.environ)
for key in list(env):
    if any(word in key.upper() for word in ('OPENAI','SMTP','TELEGRAM','GOOGLE','GMAIL')): env.pop(key)
env['ATLAS_DATA_DIR']=str(DATA)
env['PYTHONIOENCODING']='utf-8'
os.environ['PLAYWRIGHT_BROWSERS_PATH']=str(BASE/'submission/qa/video/pw-browsers')
URL='http://127.0.0.1:8877'
log=(QA/'server.log').open('w',encoding='utf-8')
server=subprocess.Popen([str(SRC/'.venv/Scripts/python.exe'),str(APP/'app.py'),'--port','8877','--data-dir',str(DATA)],cwd=APP,env=env,stdout=log,stderr=log,creationflags=subprocess.CREATE_NO_WINDOW)
for _ in range(60):
    try:
        with urlopen(URL+'/api/health', timeout=1) as res: res.read()
        break
    except Exception: time.sleep(.3)

stages=[]; current=None; started=0; outputs={}
def scene(title, narration):
    global current
    if current:
        current['end']=time.monotonic()-started
        page.screenshot(path=str(QA/(f"{len(stages):02d}-end.png")))
    current={'title':title,'narration':narration,'start':time.monotonic()-started}
    stages.append(current)
    print('SCENE',len(stages),title,flush=True)

def wait(n=1): page.wait_for_timeout(int(n*1000))
def nav(view):
    page.locator(f'.sidebar [data-view="{view}"]').first.click()
    wait(.6)
def close_run():
    if page.locator('#run-dialog').is_visible(): page.locator('#run-dialog [data-close="run-dialog"]').first.click()
def fill_form(payload, scope='#skill-fields'):
    # Resize arrays through the same buttons an operator uses, then fill visible fields.
    def arrays(value,path=''):
        if isinstance(value,dict):
            for key,val in value.items(): arrays(val, f'{path}.{key}' if path else key)
        elif isinstance(value,list):
            field=page.locator(f'{scope} fieldset[data-field-path="{path}"]')
            count=field.locator(':scope > .array-items > .array-item').count()
            while count<len(value):
                page.locator(f'{scope} [data-add-row="{path}"]').click(); count+=1
            while count>len(value):
                page.locator(f'{scope} [data-remove-row="{path}"]').last.click(); count-=1
            for i,val in enumerate(value): arrays(val,f'{path}.{i}')
    arrays(payload)
    def leaves(value,path=''):
        if isinstance(value,dict):
            for k,v in value.items(): leaves(v,f'{path}.{k}' if path else k)
        elif isinstance(value,list):
            for i,v in enumerate(value): leaves(v,f'{path}.{i}')
        else:
            el=page.locator(f'{scope} .field-value[data-field-path="{path}"]')
            if el.evaluate('(el)=>el.tagName')=='SELECT': el.select_option(str(value))
            elif isinstance(value,bool): el.set_checked(value)
            else: el.fill(str(value))
    leaves(payload)

def skill(id,payload,person='atlas', pause=1.8):
    close_run(); nav('skills')
    page.locator('#skill-search').fill(id.replace('-',' '))
    # Search uses titles/descriptions; fall back to all cards when an id is not a title.
    if not page.locator(f'[data-skill="{id}"]').count(): page.locator('#skill-search').fill('')
    page.locator(f'[data-skill="{id}"]').first.click()
    page.locator('#execute-skill-form').wait_for()
    fill_form(payload)
    page.locator('#skill-person').select_option(person)
    page.locator('#skill-dialog').evaluate('(el)=>el.scrollTop=0')
    wait(pause)
    with page.expect_response(lambda r: f'/api/skills/{id}/run' in r.url and r.request.method=='POST') as got:
        page.locator('#execute-skill').click()
    run=got.value.json()
    if run.get('status') not in ('completed','awaiting_approval'): raise RuntimeError(str(run))
    page.locator('#run-dialog').wait_for(state='visible')
    page.locator('#run-dialog h3').last.scroll_into_view_if_needed()
    wait(pause)
    outputs.setdefault(id,[]).append(run)
    return run

try:
  with sync_playwright() as pw:
    browser=pw.chromium.launch(executable_path=r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',headless=True)
    context=browser.new_context(viewport={'width':1600,'height':900},record_video_dir=str(QA/'raw'),record_video_size={'width':1600,'height':900},accept_downloads=True)
    page=context.new_page(); started=time.monotonic()
    page.goto(URL,wait_until='networkidle'); page.wait_for_selector('body[data-ready="true"]')
    scene('01 / One launch brief', 'This is Atlas Office, running locally with fictional colleagues. Our scenario is preparing the launch of a Chengdu food and tea group-tour product. We will follow the entire preparation cycle. In this credential-free recording, the operator invokes the actual skills and reviews their results. Free-form model orchestration is optional and is not shown here.')
    brief='Prepare the launch of a Chengdu food and tea group-tour product: compare a CNY 50000 pilot with a CNY 120000 full launch, assign the preparation team, plan dependencies, review the budget, and deliver a management launch-readiness pack.'
    page.locator('#chat-input').fill(brief); wait(2); page.locator('#send-message').click(); wait(2)
    nav('organization'); wait(3)

    scene('02 / Compare launch options', 'Wang, the sales lead, prepares two launch alternatives. We supply the costs, impact, risk, and weights explicitly. The executed proposal skill recommends the pilot, scoring seventy-five point seven against sixty-two for full launch. It creates an editable PowerPoint proposal and Word decision brief. The recorded result is a recommendation, not spending approval.')
    proposal=skill('proposal-package',{'title':'Chengdu tour product launch','objective':'Choose a launch approach for a Chengdu food and tea group-tour product.','audience':'Executive team','currency':'CNY','weights':{'impact':5,'risk':3,'cost':2},'options':[{'name':'Pilot first','cost':50000,'impact':4,'risk':2,'benefit':'Validate with one group.','concern':'Slower reach.'},{'name':'Full launch','cost':120000,'impact':5,'risk':4,'benefit':'Reach more customers sooner.','concern':'Higher delivery risk.'}]},'sales_manager')
    page.locator('#run-dialog').evaluate('(el)=>el.scrollTop=el.scrollHeight'); wait(3)

    scene('03 / Assign qualified colleagues', 'Operations now splits the launch work using real assigned skill identifiers and projected capacity. The result assigns planning, budget review, and management handoff to qualified colleagues. If no colleague has the required skill or hours, the same tool reports unassigned work for manager attention. These are proposed allocations; it does not secretly change saved staff capacity.')
    allocations=skill('workload-rebalance',{'tasks':[{'task':'Plan launch dependencies','required_skill':'project-plan','hours':4,'preferred_person':'product_manager'},{'task':'Review launch budget','required_skill':'budget-variance','hours':3,'preferred_person':'finance_manager'},{'task':'Prepare launch readiness pack','required_skill':'management-handoff','hours':2,'preferred_person':'director'}]},'ops_manager')
    if allocations['output'].get('unassigned'): raise RuntimeError('Expected valid assigned demo work')
    page.locator('#run-dialog').evaluate('(el)=>el.scrollTop=el.scrollHeight'); wait(2)

    scene('04 / Dependencies and accountability', 'Zhao, the product lead, converts the launch preparation into a dependency-aware schedule. The pilot review follows scope confirmation, and the executive decision follows the review. Finish dates are exclusive calendar-day boundaries. A separate responsibility matrix names the responsible specialists and makes Chen accountable for the final decision.')
    skill('project-plan',{'title':'Chengdu tour launch preparation','start':'2026-10-01','tasks':[{'id':'scope','title':'Confirm pilot scope','owner':'Zhao','days':2,'depends_on':[]},{'id':'review','title':'Pilot readiness review','owner':'He','days':3,'depends_on':['scope']},{'id':'decision','title':'Executive go-live decision','owner':'Chen','days':1,'depends_on':['review']}]},'product_manager',1)
    skill('responsibility-matrix',{'project':'Chengdu tour launch preparation','assignments':[{'task':'Confirm pilot scope','responsible':['product_manager'],'accountable':'director','consulted':['sales_manager'],'informed':['ops_manager']},{'task':'Pilot readiness review','responsible':['ops_manager'],'accountable':'director','consulted':['finance_manager'],'informed':['sales_manager']}]},'product_manager',1)

    scene('05 / Track a blocker through resolution', 'We record a launch task as blocked, with Operations as its owner. The manager can inspect and edit that saved record. For this simulation, we enter a supplied readiness outcome and mark the task done. This is an operator-recorded status transition, not automatic verification of a real supplier. The record remains available in the company register.')
    tr=skill('task-tracker',{'title':'SIMULATION - confirm pilot supplier capacity','owner':'He / Operations','due':'2026-10-06','status':'blocked'},'ops_manager')
    close_run(); nav('activity')
    page.locator('.record-row').filter(has_text='SIMULATION - confirm pilot supplier capacity').locator('[data-record]').click()
    page.locator('#record-dialog select[data-field-path="status"]').select_option('done'); wait(2)
    page.locator('#record-edit-form button.primary-btn').click(); wait(1)
    page.locator('.record-row').filter(has_text='SIMULATION - confirm pilot supplier capacity').locator('[data-record]').click(); wait(2)
    page.locator('#record-dialog [data-close="record-dialog"]').click()

    scene('06 / Budget evidence and a real review gate', 'Deng, the finance lead, calculates budget variance from supplied simulation figures. Marketing is fifteen hundred renminbi over budget. We then create a real Excel launch budget and propose a revised input value. Spreadsheet editing stops at awaiting approval. Only after review and the approval click does Atlas create a new workbook copy, preserving the original.')
    skill('budget-variance',{'currency':'CNY','items':[{'category':'Marketing','budget':10000,'actual':11500},{'category':'Pilot operations','budget':25000,'actual':24500}]},'finance_manager',1)
    book=skill('spreadsheet-create',{'title':'Chengdu launch budget - simulation','columns':['Category','Budget CNY','Forecast CNY'],'rows':[['Marketing',10000,11500],['Pilot operations',25000,24500]]},'finance_manager',1)
    from openpyxl import load_workbook
    meta=book['output']['artifacts'][0]; file_response=page.request.get(URL+meta['url']); temp=QA/'budget-preview.xlsx'; temp.write_bytes(file_response.body()); wb=load_workbook(temp); sheet=wb.sheetnames[0]; wb.close()
    edit=skill('spreadsheet-edit',{'file_id':meta['id'],'sheet':sheet,'cell':'B2','value':11500},'finance_manager',1)
    if edit['status']!='awaiting_approval': raise RuntimeError('Expected review gate')
    page.locator('#run-dialog').evaluate('(el)=>el.scrollTop=0'); wait(3)
    page.locator('[data-run-decision="approve"]').click()
    page.locator('#run-dialog .pill').filter(has_text='Completed').wait_for(timeout=10000); wait(2)

    scene('07 / Management decision and completed handoff', 'Chen records the simulated decision to proceed with pilot preparation, with the weighted comparison and reviewed budget as its rationale. The final handoff reports three preparation tasks done from supplied simulation evidence and produces a management Word brief plus an unsent email draft. The preparation loop is complete; launching a real travel product still requires authorized enterprise checks.')
    skill('decision-log',{'title':'SIMULATION - pilot preparation approved','rationale':'Pilot scores 75.7 versus 62.0. Budget input reviewed. SIM: scope, supplier capacity and readiness checklist supplied for this demonstration. No real supplier contract or spending approval.','owner':'Chen / Director','review_date':'2026-10-07'},'director',1)
    handoff=skill('management-handoff',{'title':'Chengdu launch readiness - simulated completion','audience':'Executive team','items':[{'work':'Pilot scope and option comparison','owner':'Zhao and Wang','status':'done','due':'2026-10-03','next_step':'Archive decision brief','blocker':''},{'work':'Pilot readiness and capacity check','owner':'He','status':'done','due':'2026-10-06','next_step':'Replace SIM evidence in enterprise pilot','blocker':''},{'work':'Reviewed budget and management decision','owner':'Deng and Chen','status':'done','due':'2026-10-07','next_step':'Obtain real enterprise authorization','blocker':''}]},'director',1)
    page.locator('#run-dialog').evaluate('(el)=>el.scrollTop=el.scrollHeight'); wait(2)

    scene('08 / Deliver documents and retain knowledge', 'The launch pack can be delivered as a PDF as well as Word, PowerPoint, and Excel. We save the outcome as a sourced local knowledge note, then find it with a real search. This preserves the decision and its limits for the next request. Files remain downloadable from the workspace, and the email draft remains unsent.')
    pdf=skill('pdf-create',{'title':'Chengdu tour launch - simulation outcome','sections':[{'heading':'Decision','body':'Pilot first recommended: score 75.7, compared with 62.0 for Full launch. Costs and scores are supplied simulation inputs.'},{'heading':'Preparation completed','body':'Scope, readiness review and management decision were recorded as complete for this simulation. A reviewed spreadsheet edit created a new workbook copy.'},{'heading':'Release boundary','body':'This is operator-led simulated product launch preparation. No real supplier reservation, external message, payment or product release occurred.'}]},'finance_manager',1)
    skill('knowledge-save',{'title':'Chengdu launch preparation outcome','content':'Pilot first recommended at score 75.7. SIM readiness evidence supplied. Reviewed budget copy and management handoff produced. No real product release or booking.','source':'Recorded local launch demo, 27 September 2026','owner':'Chen / Director'},'transport_coordinator',1)
    skill('knowledge-search',{'query':'Chengdu launch'},'transport_coordinator',1)
    close_run(); nav('files'); wait(3)
    artifacts=QA/'artifacts'; artifacts.mkdir(exist_ok=True)
    for f in page.request.get(URL+'/api/files').json():
        response=page.request.get(URL+f"/api/files/{f['id']}/download")
        (artifacts/f['name']).write_bytes(response.body())

    scene('09 / Inspect the execution record', 'Activity shows the actual runs across Sales, Product, Operations, Finance, IT and the Director. Thirteen different skills are demonstrated. The repository contains the source, all sixty-nine skill contracts, test evidence, and this submission. The nine-stage tourism engine is an additional domain workflow. This recording demonstrates the complete operator-led launch preparation cycle, including a real approval gate and simulated completion.')
    nav('activity'); page.mouse.wheel(0,550); wait(3); page.mouse.wheel(0,-550); wait(2)
    page.goto(URL+'/docs',wait_until='networkidle'); wait(3)
    current['end']=time.monotonic()-started
    page.screenshot(path=str(QA/'09-end.png'))
    data={'scenes':stages,'outputs':outputs,'runs':page.request.get(URL+'/api/runs').json(),'records':page.request.get(URL+'/api/records').json(),'artifacts':page.request.get(URL+'/api/files').json()}
    (QA/'recording.json').write_text(json.dumps(data,indent=2,ensure_ascii=False),encoding='utf-8')
    raw=page.video.path(); context.close(); browser.close()
    (QA/'raw-path.txt').write_text(str(raw),encoding='utf-8')
    print('RECORDED',raw,flush=True)
finally:
    server.terminate(); server.wait(timeout=10); log.close()
